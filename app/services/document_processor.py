"""Document processing service for extracting text from various file formats."""
import os
import io
import json
import logging
from pathlib import Path
from typing import Optional, Tuple, Dict, Any

import magic
import pytesseract
from PIL import Image
from PyPDF2 import PdfReader
import pdfplumber
from docx import Document as DocxDocument
from striprtf.striprtf import rtf_to_text

from app.config import get_settings

logger = logging.getLogger(__name__)
settings = get_settings()


class DocumentProcessor:
    """Process documents and extract text content."""

    SUPPORTED_MIME_TYPES = {
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "application/msword": "doc",
        "application/rtf": "rtf",
        "text/rtf": "rtf",
        "text/plain": "txt",
        "image/png": "image",
        "image/jpeg": "image",
        "image/tiff": "image",
        "image/bmp": "image",
    }

    def __init__(self, ocr_languages: Optional[str] = None):
        """Initialize processor with OCR language settings."""
        self.ocr_languages = ocr_languages or settings.ocr_languages

    def detect_mime_type(self, file_path: str) -> str:
        """Detect MIME type of a file."""
        mime = magic.Magic(mime=True)
        return mime.from_file(file_path)

    def is_supported(self, mime_type: str) -> bool:
        """Check if a MIME type is supported for processing."""
        return mime_type in self.SUPPORTED_MIME_TYPES

    def process(self, file_path: str, mime_type: Optional[str] = None) -> Dict[str, Any]:
        """
        Process a document and extract text.

        Args:
            file_path: Path to the document file
            mime_type: Optional MIME type (auto-detected if not provided)

        Returns:
            Dictionary with extracted text, page count, word count, and metadata
        """
        if mime_type is None:
            mime_type = self.detect_mime_type(file_path)

        if not self.is_supported(mime_type):
            raise ValueError(f"Unsupported file type: {mime_type}")

        file_type = self.SUPPORTED_MIME_TYPES[mime_type]

        try:
            if file_type == "pdf":
                return self._process_pdf(file_path)
            elif file_type == "docx":
                return self._process_docx(file_path)
            elif file_type == "doc":
                return self._process_doc(file_path)
            elif file_type == "rtf":
                return self._process_rtf(file_path)
            elif file_type == "txt":
                return self._process_txt(file_path)
            elif file_type == "image":
                return self._process_image(file_path)
            else:
                raise ValueError(f"No processor for file type: {file_type}")
        except Exception as e:
            logger.error(f"Error processing {file_path}: {e}")
            raise

    def _process_pdf(self, file_path: str) -> Dict[str, Any]:
        """Extract text from PDF, using OCR for scanned pages."""
        text_parts = []
        page_count = 0
        metadata = {}

        # First try pdfplumber for better text extraction
        try:
            with pdfplumber.open(file_path) as pdf:
                page_count = len(pdf.pages)
                metadata = pdf.metadata or {}

                for i, page in enumerate(pdf.pages):
                    page_text = page.extract_text() or ""

                    # If no text found, try OCR
                    if len(page_text.strip()) < 50:
                        logger.info(f"Page {i+1} appears scanned, using OCR")
                        try:
                            # Convert page to image for OCR
                            img = page.to_image(resolution=300)
                            ocr_text = pytesseract.image_to_string(
                                img.original,
                                lang=self.ocr_languages
                            )
                            page_text = ocr_text if ocr_text.strip() else page_text
                        except Exception as ocr_err:
                            logger.warning(f"OCR failed for page {i+1}: {ocr_err}")

                    text_parts.append(page_text)

        except Exception as e:
            logger.warning(f"pdfplumber failed, falling back to PyPDF2: {e}")
            # Fallback to PyPDF2
            with open(file_path, "rb") as f:
                reader = PdfReader(f)
                page_count = len(reader.pages)
                metadata = dict(reader.metadata) if reader.metadata else {}

                for page in reader.pages:
                    text_parts.append(page.extract_text() or "")

        full_text = "\n\n".join(text_parts)
        word_count = len(full_text.split())

        return {
            "extracted_text": full_text,
            "page_count": page_count,
            "word_count": word_count,
            "metadata": metadata,
        }

    def _process_docx(self, file_path: str) -> Dict[str, Any]:
        """Extract text from DOCX files."""
        doc = DocxDocument(file_path)

        # Extract all paragraphs
        paragraphs = [para.text for para in doc.paragraphs]

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)

        full_text = "\n".join(paragraphs)
        word_count = len(full_text.split())

        # Get document properties
        core_props = doc.core_properties
        metadata = {
            "author": core_props.author,
            "title": core_props.title,
            "created": str(core_props.created) if core_props.created else None,
            "modified": str(core_props.modified) if core_props.modified else None,
        }

        return {
            "extracted_text": full_text,
            "page_count": None,  # DOCX doesn't have page concept
            "word_count": word_count,
            "metadata": metadata,
        }

    def _process_doc(self, file_path: str) -> Dict[str, Any]:
        """Extract text from legacy DOC files using antiword."""
        import subprocess

        try:
            result = subprocess.run(
                ["antiword", file_path],
                capture_output=True,
                text=True,
                timeout=60,
            )
            full_text = result.stdout
        except (subprocess.SubprocessError, FileNotFoundError) as e:
            logger.error(f"antiword failed: {e}")
            raise ValueError("Could not process DOC file. antiword not available.")

        word_count = len(full_text.split())

        return {
            "extracted_text": full_text,
            "page_count": None,
            "word_count": word_count,
            "metadata": {},
        }

    def _process_rtf(self, file_path: str) -> Dict[str, Any]:
        """Extract text from RTF files."""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            rtf_content = f.read()

        full_text = rtf_to_text(rtf_content)
        word_count = len(full_text.split())

        return {
            "extracted_text": full_text,
            "page_count": None,
            "word_count": word_count,
            "metadata": {},
        }

    def _process_txt(self, file_path: str) -> Dict[str, Any]:
        """Extract text from plain text files."""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            full_text = f.read()

        word_count = len(full_text.split())

        return {
            "extracted_text": full_text,
            "page_count": None,
            "word_count": word_count,
            "metadata": {},
        }

    def _process_image(self, file_path: str) -> Dict[str, Any]:
        """Extract text from images using OCR."""
        image = Image.open(file_path)

        # Convert to RGB if necessary (for RGBA or palette images)
        if image.mode not in ("L", "RGB"):
            image = image.convert("RGB")

        full_text = pytesseract.image_to_string(image, lang=self.ocr_languages)
        word_count = len(full_text.split())

        return {
            "extracted_text": full_text,
            "page_count": 1,
            "word_count": word_count,
            "metadata": {
                "width": image.width,
                "height": image.height,
                "format": image.format,
            },
        }
